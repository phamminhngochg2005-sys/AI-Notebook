from pathlib import Path
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).parent
OUT = BASE / 'So_tay_AI_Tieu_hoc_Mockup.docx'
BLACK = '000000'
GRAY = '525A56'
FONT = 'Noto Sans'
SERIF = 'Noto Serif'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.78)
sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.82)
sec.right_margin = Inches(.82)
sec.header_distance = Inches(.34)
sec.footer_distance = Inches(.34)
sec.different_first_page_header_footer = True
WIDTH = 6.86

def font(run, size=None, bold=None, color=BLACK, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Header', 'Footer']:
    st = doc.styles[name]
    st.font.name = FONT
    st.font.color.rgb = RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), FONT)

# The stock Word template contains title borders and inherited footer tabs.
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
doc.styles['Footer'].paragraph_format.tab_stops.clear_all()

st = doc.styles['Normal']
st.font.size = Pt(11)
st.paragraph_format.line_spacing = 1.18
st.paragraph_format.space_after = Pt(7)
for name, size in [('Title',30),('Heading 1',24),('Heading 2',13),('Heading 3',11)]:
    st=doc.styles[name]
    st.font.size=Pt(size)
    st.font.bold=True
    st.paragraph_format.space_before=Pt(13 if name=='Heading 2' else 0)
    st.paragraph_format.space_after=Pt(8)
    st.paragraph_format.keep_with_next=True
    if name in ['Title','Heading 1']:
        st.font.name=SERIF
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),SERIF)

def para(text='', size=11, bold=False, color=BLACK, before=0, after=7, indent=0, style=None, keep=False):
    p=doc.add_paragraph(style=style)
    p.paragraph_format.space_before=Pt(before)
    p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.left_indent=Inches(indent)
    p.paragraph_format.keep_with_next=keep
    font(p.add_run(text),size,bold,color)
    return p

def rich(lead, text, **kw):
    p=para('',**kw)
    font(p.add_run(lead),kw.get('size',11),True)
    font(p.add_run(text),kw.get('size',11))
    return p

def h2(text):
    return doc.add_paragraph(text,'Heading 2')

def bookmark(p, name):
    s=OxmlElement('w:bookmarkStart'); s.set(qn('w:id'),str(len(doc.paragraphs)+100)); s.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),s.get(qn('w:id')))
    p._p.insert(0,s); p._p.append(e)

def hyperlink(p, text, url=None, anchor=None, size=10.3):
    link=OxmlElement('w:hyperlink')
    if url:
        rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        link.set(qn('r:id'),rid)
    if anchor: link.set(qn('w:anchor'),anchor)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    f=OxmlElement('w:rFonts'); f.set(qn('w:ascii'),FONT); f.set(qn('w:hAnsi'),FONT);pr.append(f)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));pr.append(sz)
    c=OxmlElement('w:color');c.set(qn('w:val'),'355B4B');pr.append(c)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);link.append(r);p._p.append(link)

def page(num, label, title, intro=None, anchor=None):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.line_spacing=Pt(1)
    font(p.add_run(),1).add_break(WD_BREAK.PAGE)
    para(f'{num:02d}   {label.upper()}',9,True,after=12)
    p=doc.add_paragraph(title,'Heading 1')
    if anchor: bookmark(p,anchor)
    if intro: para(intro,size=11.2,after=12)

def table(headers, rows, widths, fs=10.2):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for i,row in enumerate([headers]+rows):
        rr=t.rows[0] if i==0 else t.add_row()
        rp=rr._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');rp.append(no)
        if i==0: rp.append(OxmlElement('w:tblHeader'))
        for j,txt in enumerate(row):
            c=rr.cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,val in [('top',100),('left',120),('bottom',100),('right',120)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            if i==0 or i%2==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E2EAF0' if i==0 else 'F5F7F8');cp.append(sh)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.14
            if j==0 and widths[j]<.8:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for k,line in enumerate(str(txt).split('\n')):
                if k:font(p.add_run('\n'),fs)
                font(p.add_run(line),fs,i==0)
    para('',size=1,after=3)
    return t

header=sec.header.paragraphs[0]
header.paragraph_format.space_after=Pt(0)
font(header.add_run('SỔ TAY GIÁO VIÊN TIỂU HỌC   /   AI TRONG KẾ HOẠCH BÀI DẠY'),8,False)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.LEFT
footer.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH),WD_ALIGN_PARAGRAPH.RIGHT)
font(footer.add_run('Bản mẫu trình bày  |  Nội dung minh họa\t'),8,False,GRAY)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
first=sec.first_page_footer.paragraphs[0]
font(first.add_run('BẢN MẪU TRÌNH BÀY   •   THÁNG 10 NĂM 2026'),8.5)

# 1  Cover
para('GIÁO DỤC TIỂU HỌC',9,True,before=4,after=22)
p=doc.add_paragraph(style='Title')
font(p.add_run('Sổ tay hướng dẫn\ngiáo viên Tiểu học\nứng dụng AI'),30,True,name=SERIF)
p.paragraph_format.line_spacing=1.13
p.paragraph_format.space_after=Pt(13)
para('Trong xây dựng kế hoạch bài dạy\nphù hợp với điều kiện thực tế',16,False,after=15)
para('Dành cho giáo viên ở miền núi, vùng khó khăn\nvà lớp học còn hạn chế về thiết bị dạy học',11,after=12)
art=BASE/'assets/lop-hoc-mien-nui.png'
if art.exists():
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(13)
    p.add_run().add_picture(str(art),width=Inches(WIDTH))
    inline=p.runs[0]._r.xpath('.//wp:docPr')
    if inline:inline[0].set('descr','Minh họa ngôi trường giữa đồi núi và sách mở trên bàn giáo viên')
para('Câu lệnh mẫu • Ví dụ lớp học • Phiếu kiểm tra',10.5,True,after=5)
para('Bản minh họa để duyệt cấu trúc và trình bày.\nNội dung cần được rà soát chuyên môn trước khi hoàn thiện.',9.5,color=GRAY,after=0)

# 2  Orientation and navigation
page(2,'Cách sử dụng','Bắt đầu từ lớp học của thầy cô',
     'Sổ tay hỗ trợ giáo viên dùng AI để tạo bản dự thảo, điều chỉnh hoạt động và hoàn thiện kế hoạch bài dạy theo điều kiện đang có.',anchor='batdau')
para('Thầy cô có thể bắt đầu với một bài dạy quen thuộc. Điền bối cảnh lớp học, sao chép câu lệnh, rồi đối chiếu từng gợi ý với yêu cầu cần đạt và khả năng tổ chức thực tế.')
rich('Vai trò của AI. ','Đề xuất ý tưởng, soạn bản nháp và gợi ý phương án thay thế. Giáo viên kiểm tra độ chính xác, sự phù hợp với Chương trình GDPT 2018 [1] và quyết định nội dung đưa vào giảng dạy.')
h2('Tìm nhanh nội dung cần dùng')
toc=[('Chọn công cụ và thực hiện quy trình','congcu',3),('Điền bối cảnh lớp học','boicanh',4),('Tạo bản dự thảo bằng câu lệnh mẫu','prompt',5),('Điều chỉnh cho điều kiện thực tế','dieuchinh',6),('Xem một tiết Toán minh họa','vidu',7),('Sửa những gợi ý chưa phù hợp','suaban',8),('Kiểm tra trước khi sử dụng','kiemtra',9),('Lưu lại và dùng cho bài tiếp theo','dunglai',10)]
for text,anchor,n in toc:
    p=para('',after=9)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH-.05),WD_ALIGN_PARAGRAPH.RIGHT)
    hyperlink(p,text,anchor=anchor,size=11)
    font(p.add_run(f'\t{n:02d}'),11,True)
h2('Đọc và sao chép câu lệnh')
para('Phần trong dấu ngoặc vuông như [môn học] là chỗ thầy cô cần thay. Nếu chưa có thông tin, yêu cầu AI hỏi lại; tránh để công cụ tự điền các chi tiết quan trọng.',after=6)
para('Các trang sau là nội dung minh họa cho bản mockup, chưa phải bản sổ tay đã thẩm định chuyên môn.',9.5,color=GRAY)

# 3  Tools and workflow
page(3,'Khởi động','Chọn công cụ và làm theo 5 bước',
     'Với nhu cầu soạn bản nháp bằng tiếng Việt, thầy cô có thể bắt đầu bằng một công cụ AI hội thoại văn bản.',anchor='congcu')
rich('Ví dụ công cụ. ','ChatGPT và Gemini có thể hỗ trợ trao đổi ý tưởng và soạn thảo văn bản [3, 4]. Chọn công cụ truy cập được trên thiết bị đang có và phù hợp với quy định của nhà trường.')
para('Thử một yêu cầu ngắn, xem câu trả lời có rõ ràng và dễ sửa hay không. Kiểm tra điều kiện tài khoản, giới hạn sử dụng và quyền riêng tư ngay trên trang chính thức khi bắt đầu.')
h2('Quy trình thực hành')
for lead,text in [
    ('01  Chuẩn bị đầu vào','Xác định môn, lớp, bài học, yêu cầu cần đạt và học liệu. Ghi rõ sĩ số, thời lượng, đồ dùng và nhu cầu hỗ trợ của học sinh.'),
    ('02  Yêu cầu tạo bản nháp','Dán câu lệnh ở trang 5 cùng bối cảnh lớp học. Yêu cầu AI hỏi lại nếu thiếu thông tin.'),
    ('03  Điều chỉnh theo thực tế','Dùng các câu lệnh ở trang 6 để thay hoạt động cần thiết bị, giảm tải hoặc bổ sung hỗ trợ.'),
    ('04  Kiểm tra và chỉnh sửa','Đối chiếu kiến thức, yêu cầu cần đạt, thời lượng và cách đánh giá. Dùng phiếu ở trang 9.'),
    ('05  Lưu và rút kinh nghiệm','Lưu bản đã duyệt về thiết bị hoặc chép ra giấy. Sau tiết học, ghi lại điều cần sửa cho lần sau.')]:
    h2(lead);para(text,after=4)
rich('Khi mạng yếu. ','Dùng AI lúc có kết nối để chuẩn bị; lưu bản văn bản trước giờ lên lớp. Tiết học minh họa trong sổ tay không cần Internet hoặc máy chiếu.',size=10.5,before=9)

# 4  Classroom context form
page(4,'Phiếu chuẩn bị','Mô tả đúng điều kiện lớp học',
     'Điền phiếu này trước khi dùng câu lệnh mẫu. Có thể gõ trực tiếp vào bản Word hoặc in ra để ghi tay.',anchor='boicanh')
for heading,description in [
    ('Thông tin bài dạy','Môn và lớp: ....................................  Tên bài: ....................................\nThời lượng: ..................  Sĩ số: ..................  Ngày dạy: ..................'),
    ('Yêu cầu cần đạt','Chép hoặc tóm tắt chính xác từ tài liệu thầy cô đang sử dụng.\n........................................................................................................\nNguồn đối chiếu và trang: ....................................................................'),
    ('Đặc điểm học sinh','Kiến thức đã có, mức độ đọc hiểu, khả năng dùng tiếng Việt và nhu cầu hỗ trợ.\n........................................................................................................'),
    ('Thiết bị và học liệu','Đang có: ..........................................................................................\nKhông có hoặc không thể dùng: ..........................................................'),
    ('Điều kiện tổ chức','Không gian lớp, số nhóm, thời tiết hoặc hạn chế cần lưu ý.\n........................................................................................................'),
    ('Cách nhận biết học sinh đã hiểu','Sản phẩm, câu trả lời hoặc thao tác giáo viên sẽ quan sát.\n........................................................................................................')]:
    h2(heading);para(description,10.5,after=4)
rich('Bảo vệ thông tin học sinh. ','Chỉ mô tả theo nhóm hoặc số lượng. Không đưa tên, ảnh, điểm số gắn với danh tính hoặc thông tin riêng tư của từng em vào câu lệnh.',size=10,before=10)

# 5  Master prompt
page(5,'Câu lệnh 01','Tạo bản dự thảo kế hoạch bài dạy',
     'Thay toàn bộ phần trong dấu ngoặc vuông, rồi sao chép các đoạn bên dưới vào công cụ AI.',anchor='prompt')
prompt=[
'Bạn là trợ lý hỗ trợ giáo viên tiểu học xây dựng bản dự thảo kế hoạch bài dạy. Hãy dùng tiếng Việt rõ ràng, phù hợp lứa tuổi. Giữ nguyên yêu cầu cần đạt tôi cung cấp. Nếu thông tin thiếu hoặc mâu thuẫn, hãy hỏi lại để tôi xác minh; không tự bổ sung yêu cầu bắt buộc.',
'Thông tin bài dạy\nMôn: [môn học]. Lớp: [khối lớp]. Bài: [tên bài].\nThời lượng: [số phút]. Sĩ số: [số học sinh].\nYêu cầu cần đạt do tôi cung cấp: [nội dung và nguồn đối chiếu].\nKiến thức học sinh đã có: [mô tả ngắn].',
'Điều kiện thực tế\nĐặc điểm học sinh và nhu cầu hỗ trợ: [mô tả tổng hợp].\nThiết bị, học liệu đang có: [liệt kê].\nThiết bị không sử dụng được: [liệt kê].\nHạn chế về không gian, ngôn ngữ hoặc thời tiết: [mô tả].',
'Hãy trình bày bản dự thảo theo các phần:\n1. Yêu cầu cần đạt.\n2. Đồ dùng dạy học của giáo viên và học sinh.\n3. Các hoạt động dạy học chủ yếu.\n4. Chỗ ghi điều chỉnh sau bài dạy.',
'Với mỗi hoạt động, nêu thời gian, việc của giáo viên, việc của học sinh, sản phẩm dự kiến và cách kiểm tra. Tổng thời gian phải bằng [số phút]. Có cách hỗ trợ học sinh cần giúp đỡ và nhiệm vụ mở rộng cho học sinh hoàn thành sớm.',
'Chỉ đề xuất học liệu và hoạt động tổ chức được trong điều kiện tôi đã nêu. Không yêu cầu học sinh có điện thoại hoặc Internet nếu tôi không cung cấp điều kiện đó. Chỉ rõ phần cần giáo viên xác minh; không tự nhận bản nháp đã đạt chuẩn chương trình.'
]
for i,text in enumerate(prompt):
    p=para(text,10.7,after=10,indent=.15)
    if i in [1,2,3]:
        lines=text.split('\n');p.clear();font(p.add_run(lines[0]+'\n'),10.7,True);font(p.add_run('\n'.join(lines[1:])),10.7)
rich('Sau khi nhận kết quả. ','Giữ cuộc hội thoại để yêu cầu chỉnh tiếp. Chưa sử dụng bản nháp trên lớp khi chưa kiểm tra bằng phiếu ở trang 9.',size=10,before=4)

# 6  Adaptation prompts
page(6,'Câu lệnh 02 đến 04','Điều chỉnh theo điều kiện thực tế',
     'Dán một câu lệnh tiếp nối vào cuộc hội thoại đang có bản nháp. Thay phần trong ngoặc vuông bằng thông tin lớp mình.',anchor='dieuchinh')
h2('02  Khi thiếu thiết bị dạy học')
para('Hãy rà soát bản dự thảo trên. Lớp tôi không có [máy chiếu / máy in / Internet]. Chỉ có [bảng, phấn, thước, giấy đã dùng một mặt…]. Hãy thay các hoạt động phụ thuộc thiết bị bằng hoạt động đạt cùng mục tiêu và dùng đồ sẵn có. Nêu cách tổ chức, học liệu thay thế và cách kiểm tra kết quả. Giữ tổng thời gian [số phút].',10.7,indent=.15,after=10)
h2('03  Khi học sinh cần hỗ trợ khác nhau')
para('Trong lớp có [mô tả tổng hợp nhu cầu hỗ trợ]. Hãy giữ yêu cầu cần đạt, viết chỉ dẫn ngắn hơn và chia nhiệm vụ thành từng bước. Thêm một ví dụ làm mẫu, cách hỗ trợ theo cặp và một nhiệm vụ mở rộng cho học sinh hoàn thành sớm. Không gắn nhãn hoặc xếp loại cố định từng học sinh.',10.7,indent=.15,after=10)
h2('04  Khi hoạt động vượt thời lượng')
para('Hãy kiểm tra tổng thời gian và số nhiệm vụ trong bản dự thảo. Tôi chỉ có [số phút]. Đề xuất nội dung có thể rút gọn, giữ hoạt động cốt lõi và thời gian đánh giá cuối giờ. Trình bày phần đã sửa, lý do sửa và tổng thời gian mới; nêu điều giáo viên cần cân nhắc nếu không thể giữ đủ mục tiêu.',10.7,indent=.15,after=12)
h2('Thầy cô cần xem lại sau mỗi lần sửa')
para('Hoạt động mới có giúp học sinh đạt yêu cầu ban đầu không? Học liệu thay thế có đủ cho lớp không? Mọi học sinh có việc để làm và có cách thể hiện điều đã học không?')
rich('Ví dụ yêu cầu cụ thể. ','“Thay video mở đầu bằng 3 câu hỏi quan sát một vật thật trong lớp. Phần mở đầu tối đa 4 phút.”',size=10.5,before=8)

# 7  Worked lesson
page(7,'Ví dụ thực hành','Một tiết Toán với đồ dùng sẵn có',
     'Minh họa chủ đề chu vi hình chữ nhật ở lớp 3 trong 35 phút. Giáo viên cần đối chiếu yêu cầu cần đạt của chương trình môn học và kế hoạch đang áp dụng.',anchor='vidu')
rich('Bối cảnh giả định. ','30 học sinh; không có máy chiếu, máy in hoặc Internet trong giờ học. Có bảng, phấn, thước, bìa hình chữ nhật và vở học sinh.',size=10.5)
rich('Mục tiêu minh họa. ','Học sinh nhận biết đường bao; giải thích cách tính chu vi hình chữ nhật; giải được bài toán đơn giản với số đo cùng đơn vị.',size=10.5)
table(['Phút','Tổ chức hoạt động','Bằng chứng học tập'],[
    ['4','Khởi động\nGV đưa tấm bìa; HS chỉ đường bao và nhắc đặc điểm các cạnh đối diện.','HS chỉ đủ bốn cạnh tạo thành đường bao.'],
    ['10','Hình thành kiến thức\nHình dài 8 cm, rộng 5 cm. HS cộng bốn cạnh rồi giải thích cách viết gọn.','8 + 5 + 8 + 5 = 26 cm.\n(8 + 5) × 2 = 26 cm.'],
    ['12','Luyện tập theo cặp\nTính chu vi hình dài 6 cm, rộng 4 cm. GV hỗ trợ điền số đo từng cạnh; HS xong sớm giải thích hai cách tính.','Chu vi 20 cm. GV quan sát bài làm và hỏi cách tính.'],
    ['6','Vận dụng\nVườn dài 7 m, rộng 3 m. Tính chiều dài dây bao quanh đủ bốn cạnh. HS làm vào vở.','(7 + 3) × 2 = 20 m.\nHS nêu lời giải và đơn vị.'],
    ['3','Kiểm tra cuối giờ\nCá nhân tính chu vi hình dài 5 cm, rộng 3 cm vào vở hoặc bảng con.','Kết quả 16 cm. GV ghi nhận em cần hỗ trợ thêm.']
], [.55,3.75,2.56],fs=10)
rich('Tổng thời gian. ','4 + 10 + 12 + 6 + 3 = 35 phút.',size=10.5,after=6)
rich('Điều chỉnh sau bài dạy. ','Ghi hoạt động bị chậm, lỗi học sinh thường gặp và cách hỗ trợ trong tiết tiếp theo.',size=10.5)

# 8  Revision
page(8,'Từ bản nháp đến bài dạy','Sửa những gợi ý chưa phù hợp',
     'Đọc từng gợi ý của AI như một đề xuất cần kiểm tra. Dưới đây là hai tình huống minh họa và cách giáo viên có thể sửa.',anchor='suaban')
h2('Tình huống 1  Hoạt động cần thiết bị lớp không có')
rich('Gợi ý ban đầu. ','“Chiếu video 5 phút rồi cho từng em làm bài trên máy tính bảng.”')
rich('Giáo viên sửa. ','Dùng tấm bìa hình chữ nhật để học sinh chỉ đường bao. Học sinh làm vào vở, trao đổi theo cặp; giáo viên chọn một số bài để kiểm tra cách làm.')
rich('Cần giữ lại. ','Cơ hội quan sát, thao tác và kiểm tra mức độ hiểu. Chọn một hình thức khác phù hợp với đồ dùng và thời lượng của lớp.',size=10.5)
h2('Tình huống 2  Nội dung hoặc đáp án có sai sót')
rich('Gợi ý ban đầu. ','“Hình chữ nhật dài 8 cm, rộng 5 cm có chu vi 40 cm.”')
rich('Giáo viên sửa. ','Chu vi là 8 + 5 + 8 + 5 = 26 cm, cũng có thể tính (8 + 5) × 2 = 26 cm. Rà lại toàn bộ ví dụ và đáp án có liên quan trong bản nháp.')
rich('Cần đối chiếu. ','Phép tính, cách giải thích, đơn vị đo và mức độ phù hợp với bài đang dạy.',size=10.5)
h2('Câu lệnh yêu cầu AI tự rà soát')
para('Hãy kiểm tra bản kế hoạch vừa soạn. Liệt kê các phép tính, khái niệm, yêu cầu thiết bị và thời lượng cần xác minh. Với mỗi điểm nghi ngờ, giải thích lý do và đề xuất cách sửa. Không tự khẳng định bản kế hoạch đã chính xác hoặc được phê duyệt.',10.7,indent=.15)
rich('Quyết định của giáo viên. ','Tự giải lại bài tập và đối chiếu tài liệu chuyên môn. Việc AI kiểm tra lại chỉ hỗ trợ phát hiện vấn đề; chưa thay thế việc giáo viên xác minh.',size=10.5,before=8)

# 9  Teacher checklist
page(9,'Phiếu kiểm tra','Kiểm tra trước khi sử dụng',
     'Đánh dấu từng mục sau khi đã đọc và sửa bản dự thảo. Với mục chưa đạt, ghi rõ việc cần làm trước khi đưa vào giảng dạy.',anchor='kiemtra')
checks=[
('Yêu cầu cần đạt','Đã đối chiếu chương trình môn học, tài liệu đang dùng và kế hoạch giáo dục của nhà trường.'),
('Nội dung và đáp án','Đã kiểm tra kiến thức, dữ kiện, phép tính, lời giải và đơn vị.'),
('Đặc điểm học sinh','Nhiệm vụ vừa sức; chỉ dẫn rõ; có hỗ trợ và nhiệm vụ mở rộng phù hợp.'),
('Thiết bị và học liệu','Đồ dùng thực sự có sẵn; không có yêu cầu thiết bị ngoài điều kiện lớp học.'),
('Tổ chức và thời gian','Tổng thời gian phù hợp; nhóm học tập và không gian có thể tổ chức được.'),
('Đánh giá học sinh','Mỗi hoạt động có sản phẩm hoặc dấu hiệu quan sát; có kiểm tra cuối giờ.'),
('Thông tin và an toàn','Không lộ dữ liệu cá nhân; vật liệu và hoạt động phù hợp lứa tuổi.')
]
for title,desc in checks:
    p=para('',after=3,before=6,keep=True)
    font(p.add_run('☐  '+title),11,True)
    para(desc,10.5,after=2,indent=.22)
h2('Ghi nhận sau khi kiểm tra')
para('Điểm cần sửa: ........................................................................................\n................................................................................................................',10.5)
para('Tài liệu đã đối chiếu: ..............................................................................',10.5)
para('Quyết định:  ☐ Sẵn sàng sử dụng sau khi rà soát    ☐ Cần chỉnh tiếp',10.3)
para('Giáo viên: ..............................................  Ngày: ......................................',10.5)

# 10  Reuse and references
page(10,'Lưu và dùng lại','Chuẩn bị cho bài dạy tiếp theo',
     'Lưu câu lệnh đã hiệu quả và bản kế hoạch giáo viên đã chỉnh. Khi chuyển sang bài khác, cập nhật bối cảnh và yêu cầu cần đạt từ đầu.',anchor='dunglai')
h2('Một cách lưu đơn giản')
para('Tạo thư mục theo môn và lớp. Với mỗi bài, lưu ba tệp: câu lệnh đã dùng, kế hoạch đã rà soát và ghi chú sau tiết dạy. Dùng tên tệp có ngày để phân biệt các lần sửa.',10.7)
h2('Hai câu lệnh có thể dùng tiếp')
rich('Tạo nhiệm vụ luyện tập. ','“Dựa trên yêu cầu cần đạt [nội dung], hãy gợi ý 3 nhiệm vụ ngắn dùng [học liệu có sẵn]. Có đáp án, cách hỗ trợ và một nhiệm vụ mở rộng. Tôi sẽ kiểm tra trước khi dùng.”',size=10.5)
rich('Rút kinh nghiệm sau bài dạy. ','“Sau tiết học, tôi ghi nhận [mô tả tổng hợp, không có danh tính học sinh]. Hãy đề xuất 2 điều chỉnh cho lần dạy sau và cách kiểm tra hiệu quả của từng điều chỉnh.”',size=10.5)
h2('Tài liệu gốc để đối chiếu')
para('Khung soạn bài trong bản mẫu tham khảo hướng dẫn về kế hoạch giáo dục cấp tiểu học [2]. Giáo viên cần đối chiếu thêm chương trình môn học, sách giáo khoa và hướng dẫn đang được nhà trường áp dụng.',10,after=7)
refs=[
('[1] Bộ GDĐT  Chương trình tổng thể GDPT 2018','Văn bản chương trình trên cổng Bộ GDĐT','https://moet.gov.vn/content/tintuc/Lists/News/Attachments/8421/chuong-trinh-tong-the-ctgdpt-2018.pdf'),
('[2] Bộ GDĐT  Công văn 2345/BGDĐT-GDTH ngày 07/06/2021','Bản đăng trên cổng Trường Tiểu học Lê Quý Đôn','https://thlequydongovap.hcm.edu.vn/cong-van-cua-bo-gddt/so-2345-bgddt-gdth-vv-huong-dan-xay-dung-ke-hoach-giao-duc-cua-nha-truong-cap-t/vbct/61859/893968'),
('[3] OpenAI  Giới thiệu và cách bắt đầu với ChatGPT','Hướng dẫn chính thức trên OpenAI Help Center','https://help.openai.com/en/articles/12677804-what-is-chatgpt-faq'),
('[4] Google  Tổng quan về Gemini','Giới thiệu và giới hạn của công cụ trên trang chính thức','https://gemini.google/overview/')]
for title,desc,url in refs:
    para(title,9.5,True,after=1,before=5,keep=True)
    p=para('',9.5,after=3);hyperlink(p,desc,url=url,size=9.5)
para('Bản mẫu tháng 10/2026. Nội dung minh họa phục vụ duyệt cấu trúc và trình bày; cần rà soát chuyên môn trước khi biên soạn bản chính thức.',9,color=GRAY,before=12,after=0)

props=doc.core_properties
props.title='Sổ tay hướng dẫn giáo viên Tiểu học ứng dụng AI trong xây dựng kế hoạch bài dạy phù hợp với điều kiện thực tế'
props.subject='Bản mockup sổ tay giáo viên tiểu học'
props.author=''
props.keywords='giáo viên tiểu học, AI, kế hoạch bài dạy, bản mẫu'
props.comments=''
props.language='vi-VN'
# Set document language explicitly for spell-checking and assistive tools.
default_rpr=doc.styles['Normal']._element.get_or_add_rPr()
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'vi-VN');default_rpr.append(lang)
doc.save(OUT)
print(OUT)
