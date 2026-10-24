from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path('/Users/jason/Freelancer')
QA = BASE/'output/website-so-tay-ai/qa-v3'
OUT = BASE/'so-tay-ai-website/dist/assets/so-tay-ai-v3.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
sec.header_distance, sec.footer_distance = Inches(.25), Inches(.3)
sec.different_first_page_header_footer = True

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Noto Sans'
    st.font.color.rgb = RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans')
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Title'].paragraph_format.line_spacing = 1.06
doc.styles['Heading 1'].font.size = Pt(21)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 2'].font.size = Pt(12.5)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)

def p(text='', bold=False, size=None, color=None, after=None, before=None):
    x = doc.add_paragraph()
    r = x.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if after is not None: x.paragraph_format.space_after = Pt(after)
    if before is not None: x.paragraph_format.space_before = Pt(before)
    return x

def h(text): doc.add_heading(text, level=2)
def page(k, title):
    doc.add_page_break()
    p(f'{k:02d}  /  SỔ TAY ĐIỆN TỬ', bold=True, size=9, color='6254E8', after=8)
    doc.add_heading(title, level=1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c,w in zip(t.columns,widths): c.width = Inches(w)
    for i,x in enumerate(headers): t.rows[0].cells[i].text = x
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=x
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            b=OxmlElement('w:tcBorders')
            for e in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+e); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); b.append(el)
            pr.append(b)
            m=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'100'); e.set(qn('w:type'),'dxa'); m.append(e)
            pr.append(m)
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E4F8FC' if ri==0 else ('F7F8FC' if ri%2==0 else 'FFFFFF')); pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.line_spacing=1.05
                for r in para.runs:
                    r.font.size=Pt(10)
                    r.bold=ri==0
            if ri==0:
                trpr=row._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
    p('',after=1)
    return t

def link(para, text, url):
    rel=para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    a=OxmlElement('w:hyperlink'); a.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),'6254E8'); rp.append(c)
    r.append(rp); tx=OxmlElement('w:t'); tx.text=text; r.append(tx); a.append(r); para._p.append(a)

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Sổ tay AI  ·  Bản minh họa  ·  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('65708B')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

# 1 Cover
p('BẢN MINH HỌA ĐỂ DUYỆT', bold=True, size=9, color='6254E8', after=16)
doc.add_paragraph('Sổ tay điện tử hướng dẫn sử dụng AI hỗ trợ giáo viên Tiểu học xây dựng kế hoạch bài dạy',style='Title')
p('Dành cho giáo viên Tiểu học, hướng tới điều kiện miền núi và vùng khó khăn.',size=11.5,after=10)
x=p('',after=3); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
x.add_run().add_picture(str(QA/'cover-book.png'),width=Inches(5.9))
p('Nhóm thực hiện  Đang cập nhật\nĐơn vị  Đang cập nhật',size=10,after=10)
p('Tài liệu minh họa cấu trúc, cần rà soát chuyên môn',bold=True,size=10,after=5)
p('Phiên bản duyệt ngày 21/10/2026  •  06 trang mẫu',size=9,color='65708B')

# 2 Purpose
page(1,'Giới thiệu đề tài và phạm vi sử dụng')
p('Sổ tay hướng dẫn giáo viên dùng AI để đề xuất ý tưởng, tạo bản dự thảo và điều chỉnh kế hoạch bài dạy theo chương trình, học sinh và điều kiện thực tế. Giáo viên kiểm tra kết quả và quyết định nội dung đưa vào giảng dạy.')
h('Đối tượng và bối cảnh dự kiến')
p('Sản phẩm hướng tới giáo viên Tiểu học, đặc biệt ở trường miền núi và vùng khó khăn. Các tình huống minh họa ưu tiên lớp học có học liệu đơn giản, thiết bị dùng chung, kết nối Internet không ổn định và học sinh có mức độ sử dụng tiếng Việt khác nhau.')
p('Giáo viên có thể chuẩn bị câu lệnh và tải kết quả khi có mạng, sau đó chỉnh sửa, in hoặc sử dụng bản đã lưu trong giờ học. Các hoạt động cần có phương án thực hiện bằng bảng, giấy, vật thật và trao đổi trực tiếp.')
h('Giá trị thực hành của sổ tay')
for a,b in [
 ('Chọn công cụ','Lựa chọn công cụ truy cập được, có khả năng trao đổi bằng tiếng Việt và phù hợp điều kiện sử dụng của giáo viên.'),
 ('Dùng câu lệnh','Chọn mẫu theo môn học và nhiệm vụ, bổ sung thông tin lớp trước khi gửi AI.'),
 ('Điều chỉnh bản dự thảo','Thay học liệu, cách chia nhóm, ngôn ngữ và thời lượng theo tình hình thực tế.'),
 ('Kiểm tra trước khi dạy','Đối chiếu yêu cầu cần đạt, kiến thức, kết quả tính toán và tính khả thi của từng hoạt động.')]:
    q=p();q.add_run(a+'. ').bold=True;q.add_run(b)
h('Phạm vi bản tải mẫu')
p('Tệp này minh họa cấu trúc sổ tay và 01 câu lệnh Toán lớp 3. Thư viện 80 mẫu câu lệnh được hiển thị trên website; tệp mẫu không bao gồm toàn bộ thư viện. Nội dung chưa được thẩm định để áp dụng đại trà.')

# 3 Process
page(2,'Quy trình sử dụng trong năm bước')
for title,body,check in [
 ('1  Truy cập công cụ AI','Mở công cụ AI trên thiết bị có Internet. Tạo cuộc trò chuyện mới và xác định nhiệm vụ cần hỗ trợ, chẳng hạn lập dàn ý hoặc điều chỉnh một hoạt động.','Không nhập tên thật, hồ sơ cá nhân hoặc thông tin nhạy cảm của học sinh.'),
 ('2  Chọn câu lệnh phù hợp','Chọn mẫu trên website theo môn, lớp và mục đích. Đọc toàn bộ câu lệnh trước khi sao chép.','Dùng 01 nhiệm vụ rõ ràng cho mỗi yêu cầu; xác định dạng kết quả cần nhận.'),
 ('3  Bổ sung bối cảnh lớp học','Điền phiếu ở trang tiếp theo. Ghi yêu cầu cần đạt, năng lực tiếng Việt, sĩ số, thời lượng, học liệu và hạn chế về thiết bị.','Cung cấp yêu cầu cần đạt đã kiểm tra từ chương trình và kế hoạch dạy học của nhà trường.'),
 ('4  Tự kiểm tra kết quả','Đọc từng hoạt động; tự giải bài toán, kiểm tra đơn vị, ngôn ngữ và mức độ phù hợp với học sinh. Đánh dấu các chỗ AI suy đoán.','Nếu thiếu thông tin, yêu cầu AI hỏi lại hoặc nêu phần cần giáo viên xác minh.'),
 ('5  Điều chỉnh và hoàn thiện','Yêu cầu AI sửa đúng điểm chưa phù hợp. Giáo viên biên tập bản cuối, chuẩn bị học liệu và phương án dự phòng.','Lưu bản đã duyệt; sau tiết dạy, ghi điều chỉnh để dùng cho lần tiếp theo.')]:
    h(title);p(body,after=4);p('Điểm cần nhớ  '+check,size=10,color='65708B',after=9)
p('AI hỗ trợ soạn thảo. Giáo viên chịu trách nhiệm kiểm tra tính chính xác, sự phù hợp với Chương trình GDPT 2018 và quyết định sử dụng.',bold=True,size=10)

# 4 Context form
page(3,'Phiếu nhập bối cảnh lớp học')
p('Điền ngắn gọn trước khi dùng câu lệnh. Có thể sao chép các trường này vào cuộc trò chuyện với AI. Mô tả theo nhóm học sinh, không ghi danh tính cá nhân.',after=14)
for title,guide in [
 ('Môn học  Lớp  Tên bài','Ví dụ: Toán 3; chu vi hình chữ nhật.'),
 ('Yêu cầu cần đạt và kiến thức đã học','Trích nội dung đã được giáo viên đối chiếu; ghi kiến thức học sinh đã nắm.'),
 ('Sĩ số và cách tổ chức lớp','Tổng số học sinh; lớp đơn hoặc ghép; số nhóm và khả năng hỗ trợ từng nhóm.'),
 ('Năng lực sử dụng tiếng Việt của học sinh','Mức độ nghe hiểu chỉ dẫn, đọc đề, diễn đạt; nhu cầu dùng hình ảnh, vật thật hoặc hỗ trợ ngôn ngữ.'),
 ('Học liệu và thiết bị hiện có','Bảng, phấn, giấy, thước, vật thật; số lượng thiết bị thực sự dùng được.'),
 ('Điều kiện thực hiện','Điện, mạng, không gian lớp, chỗ ngồi; khó khăn dự kiến và vật liệu thay thế.'),
 ('Thời lượng và cách đánh giá','Số phút; sản phẩm mong đợi; dấu hiệu quan sát để biết học sinh đã hiểu.')]:
    p(title,bold=True,size=11,after=2)
    p(guide,size=9,color='65708B',after=2)
    p('..................................................................................................................',size=11,color='B8BED0',after=10)
p('Gợi ý nhập thêm  “Không sử dụng máy chiếu. Mỗi nhóm có một tờ giấy và một thước. Dùng câu ngắn, minh họa trực quan và có phương án cho học sinh cần hỗ trợ.”',size=10)

# 5 Prompt and plan
page(4,'Ví dụ Toán lớp 3 về chu vi')
p('Bối cảnh minh họa  28 học sinh, 7 nhóm; một số em cần hỗ trợ đọc hiểu tiếng Việt; có bảng, phấn, giấy và thước; không có máy chiếu. Thời lượng 35 phút.',size=10,after=8)
h('Câu lệnh mẫu có thể sao chép')
p('Bạn là trợ lý hỗ trợ giáo viên Tiểu học. Hãy dự thảo kế hoạch Toán lớp 3 về chu vi hình chữ nhật trong 35 phút theo bối cảnh trên. Mục tiêu dự kiến: học sinh nhận biết chu vi là độ dài đường bao, tính được chu vi hình chữ nhật với số đo cùng đơn vị và giải thích cách làm bằng lời hoặc hình vẽ. Giáo viên sẽ đối chiếu yêu cầu cần đạt trước khi sử dụng.\n\nHãy thiết kế 4 hoạt động với tổng thời gian đúng 35 phút. Mỗi hoạt động nêu việc giáo viên làm, nhiệm vụ học sinh và dấu hiệu đánh giá. Dùng câu ngắn, vật thật và học liệu sẵn có; không yêu cầu Internet hoặc máy chiếu. Có hỗ trợ học sinh đọc đề còn chậm và nhiệm vụ mở rộng. Dùng ví dụ hình chữ nhật dài 6 cm, rộng 4 cm; tự kiểm tra kết quả. Nêu các giả định cần giáo viên xác minh.',size=10,after=8)
h('Bản kế hoạch minh họa để rà soát')
table(['Hoạt động','Phút','Tổ chức và dấu hiệu đánh giá'],[
 ('Khởi động','5','Chỉ đường bao một quyển vở. Học sinh lần theo mép và nói “đường bao”; giáo viên kiểm tra cách hiểu.'),
 ('Hình thành kiến thức','12','Nhóm vẽ hoặc đo hình 6 cm × 4 cm, cộng bốn cạnh; thảo luận cách tính. Học sinh chỉ đúng các cạnh bằng nhau.'),
 ('Luyện tập','12','Cặp đôi tính chu vi hình 5 cm × 3 cm và giải thích. Đáp số 16 cm. Học sinh đọc chậm dùng hình có ghi số đo.'),
 ('Vận dụng và đánh giá','6','Học sinh nêu cách tìm chu vi mặt bàn bằng đo đạc. Phiếu ra về: tính chu vi hình 7 cm × 2 cm, đáp số 18 cm.')
],[1.45,.5,4.95])
q=p('Kiểm tra ví dụ chính  ',bold=True,size=10,after=5)
math=OxmlElement('m:oMath'); mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text='(6 + 4) × 2 = 20 cm';mr.append(mt);math.append(mr);q._p.append(math)
p('Tổng thời gian  5 + 12 + 12 + 6 = 35 phút. Mở rộng: tìm một hình chữ nhật khác có chu vi 20 cm và giải thích.',size=10,after=3)

# 6 Review and reference
page(5,'Kiểm tra và hoàn thiện trước khi dạy')
p('Đánh dấu sau khi giáo viên đã kiểm tra. Với mục chưa đạt, sửa bản kế hoạch hoặc hỏi tổ chuyên môn trước khi áp dụng.',after=11)
for a,b in [
 ('Yêu cầu cần đạt','Đã đối chiếu với Chương trình GDPT 2018, nội dung môn học và kế hoạch của nhà trường.'),
 ('Kiến thức và phép tính','Đã tự giải bài, kiểm tra công thức, số liệu, thuật ngữ và đơn vị; ví dụ chính có chu vi 20 cm.'),
 ('Mức độ phù hợp với học sinh','Chỉ dẫn dễ hiểu, có hình hoặc vật thật, có hỗ trợ tiếng Việt và nhiệm vụ phù hợp các mức độ.'),
 ('Học liệu và khả năng triển khai','Có đủ vật liệu, không phụ thuộc thiết bị chưa có; đã chuẩn bị phương án dự phòng.'),
 ('Tiến trình và đánh giá','Tổng thời gian đúng 35 phút; nhiệm vụ, sản phẩm và dấu hiệu đánh giá quan sát được.'),
 ('Bảo mật và trách nhiệm','Không đưa thông tin cá nhân học sinh lên AI; giáo viên đã đọc và duyệt toàn bộ bản cuối.')]:
    q=p(after=9);q.add_run('☐  '+a+'. ').bold=True;q.add_run(b)
h('Định hướng tham khảo Công văn 2345')
p('Tham khảo Công văn số 2345/BGDĐT-GDTH ngày 07/06/2021 của Bộ Giáo dục và Đào tạo về hướng dẫn xây dựng kế hoạch giáo dục của nhà trường cấp tiểu học. Khi hoàn thiện, rà soát các thành phần của kế hoạch bài dạy: yêu cầu cần đạt, đồ dùng dạy học, hoạt động dạy học chủ yếu và điều chỉnh sau bài dạy nếu có. Giáo viên tiếp tục đối chiếu hướng dẫn hiện hành của địa phương và nhà trường.',size=10)
q=p(size=9,after=12);link(q,'Thông tin văn bản trên cổng Bộ Giáo dục và Đào tạo','https://moet-edu.coquan.net/giaoducquocdan/giao-duc-tieu-hoc/Pages/chi-tiet-van-ban-chi-dao-dieu-hanh.aspx%3FItemID=2967')
h('Ghi chú sau tiết dạy')
p('Điều phù hợp  ..............................................................................................\nĐiều cần thay đổi  ........................................................................................\nHọc liệu cần bổ sung  ...................................................................................',size=10,after=13)
p('Tài liệu minh họa cấu trúc, cần rà soát chuyên môn',bold=True,size=10,after=5)
p('Thư viện 80 mẫu câu lệnh được hiển thị trên website. Bản tải này có 01 câu lệnh minh họa và không thay thế bản sổ tay hoàn chỉnh.',size=9,color='65708B')

doc.core_properties.title='Sổ tay điện tử hướng dẫn sử dụng AI hỗ trợ giáo viên Tiểu học xây dựng kế hoạch bài dạy'
doc.core_properties.subject='Bản minh họa sáu trang để duyệt cấu trúc và giao diện'
doc.core_properties.author='Đang cập nhật'
doc.core_properties.keywords='AI, giáo viên Tiểu học, kế hoạch bài dạy, minh họa'
for root in [doc._element, doc.styles.element]:
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for el in doc._element.iter(qn('wp:docPr')):
    el.set('descr','Sách mở với bìa tím indigo và các nút kết nối AI màu tím và cyan')
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
